#!/usr/bin/env python3
"""Extract the MineVision Career staging dataset from a structured DOCX file."""

from __future__ import annotations

import argparse
import json
import os
import re
import stat
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


SCHEMA_VERSION = "1.0"
REQUIRED_CATEGORY_SLUGS = (
    "eksplorasi-dan-geologi",
    "perencanaan-dan-rekayasa-tambang",
    "operasi-dan-produksi-tambang",
    "survei-gis-dan-geospasial",
    "pengolahan-mineral-dan-metalurgi",
    "keselamatan-kesehatan-dan-tanggap-darurat",
    "lingkungan-dan-keberlanjutan",
    "pemeliharaan-dan-keandalan-peralatan",
    "logistik-rantai-pasok-dan-pengadaan",
    "data-teknologi-informasi-dan-otomasi",
    "keuangan-komersial-dan-manajemen",
    "legal-perizinan-dan-kepatuhan",
    "sumber-daya-manusia-dan-administrasi",
)
REQUIRED_PROFILE_SECTIONS = (
    "work_scope",
    "competency",
    "education",
    "software",
    "training",
)
CATEGORY_NAMES = {
    "eksplorasi-dan-geologi": "Eksplorasi dan Geologi",
    "perencanaan-dan-rekayasa-tambang": "Perencanaan dan Rekayasa Tambang",
    "operasi-dan-produksi-tambang": "Operasi dan Produksi Tambang",
    "survei-gis-dan-geospasial": "Survei, GIS, dan Geospasial",
    "pengolahan-mineral-dan-metalurgi": "Pengolahan Mineral dan Metalurgi",
    "keselamatan-kesehatan-dan-tanggap-darurat": (
        "Keselamatan, Kesehatan, dan Tanggap Darurat"
    ),
    "lingkungan-dan-keberlanjutan": "Lingkungan dan Keberlanjutan",
    "pemeliharaan-dan-keandalan-peralatan": (
        "Pemeliharaan dan Keandalan Peralatan"
    ),
    "logistik-rantai-pasok-dan-pengadaan": (
        "Logistik, Rantai Pasok, dan Pengadaan"
    ),
    "data-teknologi-informasi-dan-otomasi": (
        "Data, Teknologi Informasi, dan Otomasi"
    ),
    "keuangan-komersial-dan-manajemen": (
        "Keuangan, Komersial, dan Manajemen"
    ),
    "legal-perizinan-dan-kepatuhan": "Legal, Perizinan, dan Kepatuhan",
    "sumber-daya-manusia-dan-administrasi": (
        "Sumber Daya Manusia dan Administrasi"
    ),
}
SECTION_ALIASES = {
    "work_scope": {
        "ruang lingkup",
        "ruang lingkup kerja",
        "ruang lingkup pekerjaan",
        "tugas dan tanggung jawab",
        "lingkup pekerjaan",
    },
    "competency": {
        "kompetensi",
        "kompetensi utama",
        "keterampilan",
        "keterampilan dan kompetensi",
    },
    "education": {
        "pendidikan",
        "pendidikan yang relevan",
        "latar belakang pendidikan",
        "kualifikasi pendidikan",
    },
    "software": {
        "software",
        "software yang digunakan",
        "perangkat lunak",
        "software dan teknologi",
        "perangkat lunak dan teknologi",
    },
    "training": {
        "pelatihan",
        "pelatihan relevan",
        "training",
        "sertifikasi",
        "pelatihan dan sertifikasi",
        "training dan sertifikasi",
    },
}
PROFESSION_ALIASES = {
    "profesi",
    "daftar profesi",
    "profesi utama",
    "jenis profesi",
    "jabatan",
    "posisi pekerjaan",
}
LIST_PREFIX_RE = re.compile(
    r"^\s*(?:(?:\d+|[A-Za-z]|[IVXLCDM]+)[.)]\s+|[-*•▪◦]\s*)",
    re.IGNORECASE,
)
RAW_URL_RE = re.compile(r"https://[^\s<>()\[\]{}]+", re.IGNORECASE)


@dataclass(frozen=True)
class Block:
    kind: str
    text: str = ""
    style: str = ""
    heading_level: int | None = None
    list_number_id: int | None = None
    rows: tuple[tuple[str, ...], ...] = ()
    urls: tuple[str, ...] = ()


@dataclass
class ParseReport:
    paragraph_count: int = 0
    table_count: int = 0
    software_table_count: int = 0
    software_tables_by_category: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def normalize_label(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", normalize_space(value))
    normalized = "".join(
        character for character in normalized if not unicodedata.combining(character)
    )
    normalized = normalized.casefold()
    normalized = re.sub(r"^kategori\s+profesi\s+", "", normalized)
    normalized = re.sub(r"^kategori\s+(?:ke-?)?\d+\s*[:.)-]?\s*", "", normalized)
    normalized = re.sub(
        r"^(?:(?:\d+|[a-z]|[ivxlcdm]+)[.)-]\s*)+",
        "",
        normalized,
    )
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized)
    return normalize_space(normalized)


def slugify(value: str, max_length: int) -> str:
    slug = normalize_label(value).replace(" ", "-")
    slug = re.sub(r"-+", "-", slug).strip("-")
    if len(slug) <= max_length:
        return slug
    return slug[:max_length].rstrip("-")


def strip_list_prefix(value: str) -> str:
    return normalize_space(LIST_PREFIX_RE.sub("", value, count=1))


def heading_level(paragraph: Paragraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    match = re.search(r"(?:Heading|Judul)\s+(\d+)$", style_name, re.IGNORECASE)
    if match:
        return int(match.group(1))
    return None


def list_number_id(paragraph: Paragraph) -> int | None:
    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    if numbering is None or numbering.numId is None:
        return None
    return int(numbering.numId.val)


def paragraph_urls(paragraph: Paragraph) -> tuple[str, ...]:
    urls: list[str] = []
    relationships = paragraph.part.rels

    for hyperlink in paragraph._p.iter(qn("w:hyperlink")):
        relationship_id = hyperlink.get(qn("r:id"))
        if relationship_id and relationship_id in relationships:
            target = relationships[relationship_id].target_ref
            if target.startswith("https://"):
                urls.append(target)

    urls.extend(RAW_URL_RE.findall(paragraph.text))
    return tuple(dict.fromkeys(url.rstrip(".,;:") for url in urls))


def cell_text(cell: Any) -> str:
    return normalize_space("\n".join(paragraph.text for paragraph in cell.paragraphs))


def table_urls(table: Table) -> tuple[str, ...]:
    urls: list[str] = []
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                urls.extend(paragraph_urls(paragraph))
    return tuple(dict.fromkeys(urls))


def read_blocks(document: DocumentObject) -> list[Block]:
    blocks: list[Block] = []

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = normalize_space(item.text)
            if not text:
                continue
            blocks.append(
                Block(
                    kind="paragraph",
                    text=text,
                    style=item.style.name if item.style is not None else "",
                    heading_level=heading_level(item),
                    list_number_id=list_number_id(item),
                    urls=paragraph_urls(item),
                )
            )
            continue

        if isinstance(item, Table):
            rows = tuple(
                tuple(cell_text(cell) for cell in row.cells) for row in item.rows
            )
            blocks.append(
                Block(kind="table", rows=rows, urls=table_urls(item))
            )

    return blocks


def category_slug_from_heading(value: str) -> str | None:
    normalized = normalize_label(value)
    for slug, name in CATEGORY_NAMES.items():
        candidates = {
            normalize_label(name),
            normalize_label(slug.replace("-", " ")),
        }
        if normalized in candidates:
            return slug
    return None


def section_from_heading(value: str) -> str | None:
    normalized = normalize_label(value)
    for section, aliases in SECTION_ALIASES.items():
        if normalized in aliases:
            return section
    return None


def is_profession_heading(value: str) -> bool:
    return normalize_label(value) in PROFESSION_ALIASES | {"contoh profesi"}


def is_probable_subheading(block: Block) -> bool:
    if block.kind != "paragraph":
        return False
    if block.heading_level is not None:
        return True
    text = strip_list_prefix(block.text)
    return len(text) <= 90 and text.endswith(":")


def markdown_escape(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", "<br>")


def table_to_markdown(rows: Sequence[Sequence[str]]) -> str:
    cleaned_rows = [list(row) for row in rows if any(normalize_space(cell) for cell in row)]
    if not cleaned_rows:
        return ""
    width = max(len(row) for row in cleaned_rows)
    normalized_rows = [row + [""] * (width - len(row)) for row in cleaned_rows]
    header = normalized_rows[0]
    body = normalized_rows[1:]
    lines = [
        "| " + " | ".join(markdown_escape(cell) for cell in header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend(
        "| " + " | ".join(markdown_escape(cell) for cell in row) + " |"
        for row in body
    )
    return "\n".join(lines)


def blocks_to_markdown(blocks: Sequence[Block]) -> str:
    output: list[str] = []
    for block in blocks:
        if block.kind == "table":
            markdown = table_to_markdown(block.rows)
            if markdown:
                output.append(markdown)
            continue

        text = block.text
        if block.heading_level is not None:
            level = min(max(block.heading_level, 2), 4)
            output.append(f"{'#' * level} {text}")
        elif LIST_PREFIX_RE.match(text):
            output.append(f"- {strip_list_prefix(text)}")
        else:
            output.append(text)
    return "\n\n".join(output).strip()


def split_document(
    blocks: Sequence[Block], report: ParseReport
) -> tuple[dict[str, list[Block]], list[Block]]:
    reference_index: int | None = None
    for index, block in enumerate(blocks):
        if block.kind == "paragraph" and block.text.strip() == "REFERENSI":
            reference_index = index
            break

    if reference_index is None:
        report.errors.append('Heading persis "REFERENSI" tidak ditemukan.')
        material_blocks = list(blocks)
        reference_blocks: list[Block] = []
    else:
        material_blocks = list(blocks[:reference_index])
        reference_blocks = list(blocks[reference_index + 1 :])

    categories: dict[str, list[Block]] = {}
    current_slug: str | None = None
    found_order: list[str] = []

    for block in material_blocks:
        if block.kind == "paragraph":
            category_slug = category_slug_from_heading(block.text)
            if category_slug is not None:
                if category_slug in categories:
                    report.errors.append(
                        f"Heading kategori materi muncul lebih dari sekali: {category_slug}."
                    )
                else:
                    categories[category_slug] = []
                    found_order.append(category_slug)
                current_slug = category_slug
                continue

        if current_slug is not None:
            categories[current_slug].append(block)

    missing_categories = [
        slug for slug in REQUIRED_CATEGORY_SLUGS if slug not in categories
    ]
    unexpected_categories = [
        slug for slug in categories if slug not in REQUIRED_CATEGORY_SLUGS
    ]
    if missing_categories:
        report.errors.append(
            "Kategori materi wajib tidak ditemukan: " + ", ".join(missing_categories)
        )
    if unexpected_categories:
        report.errors.append(
            "Kategori materi tidak dikenal: " + ", ".join(unexpected_categories)
        )
    if found_order != list(REQUIRED_CATEGORY_SLUGS):
        report.errors.append(
            "Urutan heading kategori materi tidak sesuai urutan 13 kategori wajib."
        )

    return categories, reference_blocks


def unique_key(base: str, used: set[str], max_length: int) -> str:
    candidate = slugify(base, max_length) or "item"
    if candidate not in used:
        used.add(candidate)
        return candidate
    suffix = 2
    while True:
        suffix_text = f"-{suffix}"
        shortened = candidate[: max_length - len(suffix_text)].rstrip("-")
        resolved = f"{shortened}{suffix_text}"
        if resolved not in used:
            used.add(resolved)
            return resolved
        suffix += 1


def split_cell_values(value: str) -> list[str]:
    values = [normalize_space(part) for part in re.split(r"[\n;•]+", value)]
    return [strip_list_prefix(value) for value in values if strip_list_prefix(value)]


def table_header_indexes(rows: Sequence[Sequence[str]]) -> dict[str, int]:
    if not rows:
        return {}
    indexes: dict[str, int] = {}
    for index, header in enumerate(rows[0]):
        normalized = normalize_label(header)
        if any(token in normalized for token in ("profesi", "jabatan", "posisi", "nama")):
            indexes.setdefault("name", index)
        if "kelompok" in normalized or normalized in {"group", "grup"}:
            indexes.setdefault("group", index)
        if "deskripsi" in normalized or "keterangan" in normalized:
            indexes.setdefault("description", index)
        if "kategori" in normalized:
            indexes.setdefault("category", index)
        if "sumber" in normalized or "referensi" in normalized or "judul" in normalized:
            indexes.setdefault("source", index)
        if "organisasi" in normalized or "institusi" in normalized or "penerbit" in normalized:
            indexes.setdefault("organization", index)
        if "url" in normalized or "tautan" in normalized or "link" in normalized:
            indexes.setdefault("url", index)
    return indexes


def extract_professions(
    category_slug: str,
    blocks: Sequence[Block],
    report: ParseReport,
) -> list[dict[str, Any]]:
    professions: list[dict[str, Any]] = []
    used_natural_keys: set[str] = set()
    group_orders: dict[str, int] = {}
    current_mode: str | None = None
    current_group_label = "Profesi Utama"
    group_heading_number_id: int | None = None

    def add_profession(
        name: str,
        group_label: str,
        description: str | None = None,
    ) -> None:
        clean_name = strip_list_prefix(name)
        clean_group = strip_list_prefix(group_label) or "Profesi Utama"
        if not clean_name or len(clean_name) > 180:
            if clean_name:
                report.warnings.append(
                    f"{category_slug}: kandidat profesi dilewati karena lebih dari 180 karakter: {clean_name[:80]}..."
                )
            return
        group_key = slugify(clean_group, 120)
        profession_slug = slugify(clean_name, 200)
        if not group_key or not profession_slug:
            return
        natural_key = f"{group_key}\0{profession_slug}"
        if natural_key in used_natural_keys:
            report.warnings.append(
                f"{category_slug}: profesi duplikat dilewati: {clean_name}."
            )
            return
        used_natural_keys.add(natural_key)
        display_order = group_orders.get(group_key, 0)
        group_orders[group_key] = display_order + 1
        professions.append(
            {
                "groupKey": group_key,
                "groupLabel": clean_group,
                "name": clean_name,
                "slug": profession_slug,
                "description": normalize_space(description) if description else None,
                "displayOrder": display_order,
            }
        )

    for block in blocks:
        if block.kind == "paragraph":
            section = section_from_heading(block.text)
            if section is not None:
                current_mode = section
                continue
            if is_profession_heading(block.text):
                current_mode = "professions"
                current_group_label = "Profesi Utama"
                group_heading_number_id = None
                continue
            if current_mode == "professions":
                # Dokumen sumber memakai list terpisah untuk label kelompok dan
                # nama profesi. Paragraf biasa tepat setelah "Contoh Profesi"
                # adalah catatan editorial, bukan nama profesi.
                if block.list_number_id is None:
                    continue
                if group_heading_number_id is None:
                    group_heading_number_id = block.list_number_id
                    current_group_label = strip_list_prefix(block.text).rstrip(":")
                    continue
                if block.list_number_id == group_heading_number_id:
                    current_group_label = strip_list_prefix(block.text).rstrip(":")
                    continue
                add_profession(block.text, current_group_label)
            continue

        if current_mode != "professions":
            continue
        rows = [row for row in block.rows if any(normalize_space(cell) for cell in row)]
        if not rows:
            continue
        indexes = table_header_indexes(rows)
        data_rows = rows[1:] if indexes else rows
        for row in data_rows:
            name_index = indexes.get("name", 0)
            if name_index >= len(row):
                continue
            group_label = current_group_label
            group_index = indexes.get("group")
            if group_index is not None and group_index < len(row) and row[group_index]:
                group_label = row[group_index]
            description = None
            description_index = indexes.get("description")
            if (
                description_index is not None
                and description_index < len(row)
                and row[description_index]
            ):
                description = row[description_index]
            for name in split_cell_values(row[name_index]):
                add_profession(name, group_label, description)

    return professions


def extract_profile_items(
    category_slug: str,
    blocks: Sequence[Block],
    report: ParseReport,
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    used_keys: set[str] = set()
    group_orders: dict[tuple[str, str], int] = {}
    current_section: str | None = None
    current_group_label = "Informasi Utama"

    def add_item(value: str, group_label: str) -> None:
        if current_section is None:
            return
        clean_value = strip_list_prefix(value)
        clean_group = strip_list_prefix(group_label) or "Informasi Utama"
        if not clean_value:
            return
        group_key = slugify(clean_group, 120)
        if not group_key:
            return
        identity = (current_section, group_key)
        display_order = group_orders.get(identity, 0)
        group_orders[identity] = display_order + 1
        key_base = f"{current_section}-{group_key}-{clean_value}"
        item_key = unique_key(key_base, used_keys, 160)
        items.append(
            {
                "itemKey": item_key,
                "section": current_section,
                "groupKey": group_key,
                "groupLabel": clean_group,
                "value": clean_value,
                "displayOrder": display_order,
            }
        )

    for block in blocks:
        if block.kind == "paragraph":
            section = section_from_heading(block.text)
            if section is not None:
                current_section = section
                current_group_label = "Informasi Utama"
                continue
            if is_profession_heading(block.text):
                current_section = None
                continue
            if current_section is None:
                continue
            if is_probable_subheading(block) and not LIST_PREFIX_RE.match(block.text):
                current_group_label = block.text.rstrip(":")
            else:
                add_item(block.text, current_group_label)
            continue

        if current_section is None:
            continue
        if current_section == "software":
            report.software_table_count += 1
            report.software_tables_by_category[category_slug] = (
                report.software_tables_by_category.get(category_slug, 0) + 1
            )
        rows = [row for row in block.rows if any(normalize_space(cell) for cell in row)]
        if not rows:
            continue
        headers = list(rows[0])
        data_rows = rows[1:] if len(rows) > 1 else rows
        for row in data_rows:
            populated = [
                (headers[index] if index < len(headers) else "", cell)
                for index, cell in enumerate(row)
                if normalize_space(cell)
            ]
            if not populated:
                continue
            group_label = current_group_label
            first_header = normalize_label(populated[0][0])
            if first_header in {"kelompok", "group", "grup", "kategori"}:
                group_label = populated[0][1]
                populated = populated[1:]
            value_parts = [
                f"{header}: {cell}" if normalize_space(header) else cell
                for header, cell in populated
            ]
            add_item("; ".join(value_parts), group_label)

    present_sections = {item["section"] for item in items}
    for section in REQUIRED_PROFILE_SECTIONS:
        if section not in present_sections:
            report.errors.append(
                f"{category_slug}: section profile item tidak memiliki data: {section}."
            )
    return items


def first_general_paragraph(blocks: Sequence[Block]) -> str | None:
    for block in blocks:
        if block.kind != "paragraph":
            continue
        if section_from_heading(block.text) is not None or is_profession_heading(block.text):
            continue
        if is_probable_subheading(block):
            continue
        text = strip_list_prefix(block.text)
        if text:
            return text
    return None


def clean_reference_text(text: str, urls: Iterable[str]) -> str:
    cleaned = text
    for url in urls:
        cleaned = cleaned.replace(url, " ")
    return strip_list_prefix(normalize_space(cleaned).strip(" -–—|"))


def source_name_from_url(url: str) -> str:
    return re.sub(r"^https://", "", url, flags=re.IGNORECASE).split("/", 1)[0]


def infer_source_type(text: str) -> str:
    normalized = normalize_label(text)
    if any(token in normalized for token in ("undang undang", "peraturan", "regulation")):
        return "regulation"
    if any(token in normalized for token in ("badan pusat statistik", "statistics agency")):
        return "statistics_agency"
    if any(token in normalized for token in ("annual report", "laporan tahunan", "sustainability report")):
        return "company_report"
    if any(token in normalized for token in ("jurnal", "journal", "universitas", "university", "academic")):
        return "academic"
    if any(token in normalized for token in ("kementerian", "pemerintah", "esdm", "bappenas")):
        return "government"
    if any(token in normalized for token in ("harga", "market data", "bursa")):
        return "market_data"
    return "other"


def infer_organization(text: str, url: str | None) -> str:
    for separator in (". ", " – ", " — ", " - ", ": "):
        candidate = normalize_space(text.split(separator, 1)[0])
        if 1 <= len(candidate) <= 200:
            return candidate
    if 1 <= len(text) <= 200:
        return text
    if url:
        hostname = re.sub(r"^https://", "", url, flags=re.IGNORECASE).split("/", 1)[0]
        return hostname[:200]
    return ""


def extract_reference_records(block: Block) -> list[tuple[str, str | None]]:
    if block.kind == "paragraph":
        urls = list(block.urls)
        clean_text = clean_reference_text(block.text, urls)
        if clean_text:
            return [(clean_text, urls[0] if urls else None)]
        return [(source_name_from_url(url), url) for url in urls]

    rows = [row for row in block.rows if any(normalize_space(cell) for cell in row)]
    if not rows:
        return []
    indexes = table_header_indexes(rows)
    data_rows = rows[1:] if indexes else rows
    records: list[tuple[str, str | None]] = []
    for row in data_rows:
        row_text = " — ".join(cell for cell in row if normalize_space(cell))
        row_urls = RAW_URL_RE.findall(row_text)
        clean_text = clean_reference_text(row_text, row_urls)
        if clean_text:
            records.append((clean_text, row_urls[0] if row_urls else None))
    return records


def parse_references(
    blocks: Sequence[Block], report: ParseReport
) -> tuple[list[dict[str, Any]], dict[str, list[str]]]:
    source_catalog: list[dict[str, Any]] = []
    source_slugs_by_category = {slug: [] for slug in REQUIRED_CATEGORY_SLUGS}
    source_slug_by_identity: dict[tuple[str, str], str] = {}
    used_source_slugs: set[str] = set()
    current_slug: str | None = None

    def add_source(category_slug: str, text: str, url: str | None) -> None:
        name = normalize_space(text)
        if not name:
            return
        if len(name) > 200:
            report.errors.append(
                f"{category_slug}: nama sumber melebihi 200 karakter dan tidak dipotong otomatis."
            )
            return
        organization = infer_organization(name, url)
        if not organization:
            report.errors.append(
                f"{category_slug}: organisasi sumber tidak dapat ditentukan dari dokumen."
            )
            return
        identity = (normalize_label(name), url or "")
        source_slug = source_slug_by_identity.get(identity)
        if source_slug is None:
            source_slug = unique_key(organization or name, used_source_slugs, 220)
            source_type = infer_source_type(f"{organization} {name}")
            source_catalog.append(
                {
                    "slug": source_slug,
                    "name": name,
                    "type": source_type,
                    "organization": organization,
                    "url": url,
                    "description": None,
                    "isOfficial": source_type
                    in {"government", "statistics_agency", "company_report", "regulation"},
                    "verificationStatus": "verified",
                }
            )
            source_slug_by_identity[identity] = source_slug
        if source_slug not in source_slugs_by_category[category_slug]:
            source_slugs_by_category[category_slug].append(source_slug)

    for block in blocks:
        if block.kind == "paragraph":
            category_slug = category_slug_from_heading(block.text)
            if category_slug is not None:
                current_slug = category_slug
                continue

        if current_slug is not None:
            for text, url in extract_reference_records(block):
                add_source(current_slug, text, url)
            continue

        if block.kind == "table":
            rows = [row for row in block.rows if any(normalize_space(cell) for cell in row)]
            indexes = table_header_indexes(rows)
            category_index = indexes.get("category")
            source_index = indexes.get("source")
            if category_index is None or source_index is None:
                report.warnings.append(
                    "Tabel referensi sebelum heading kategori dilewati karena tidak memiliki kolom kategori dan sumber."
                )
                continue
            url_index = indexes.get("url")
            for row in rows[1:]:
                if category_index >= len(row) or source_index >= len(row):
                    continue
                category_slug = category_slug_from_heading(row[category_index])
                if category_slug is None:
                    report.warnings.append(
                        f"Kategori pada tabel referensi tidak dikenali: {row[category_index]}."
                    )
                    continue
                url = None
                if url_index is not None and url_index < len(row):
                    matches = RAW_URL_RE.findall(row[url_index])
                    url = matches[0].rstrip(".,;:") if matches else None
                add_source(category_slug, row[source_index], url)
        elif block.text:
            report.warnings.append(
                f"Referensi di luar kelompok kategori dilewati: {block.text[:100]}."
            )

    for category_slug, source_slugs in source_slugs_by_category.items():
        if not source_slugs:
            report.errors.append(
                f"{category_slug}: tidak memiliki sumber pada bagian REFERENSI."
            )

    return source_catalog, source_slugs_by_category


def build_dataset(
    document: DocumentObject,
    input_path: Path,
    effective_date: str,
    published_at: str,
) -> tuple[dict[str, Any], dict[str, dict[str, Any]], ParseReport]:
    report = ParseReport(
        paragraph_count=len(document.paragraphs),
        table_count=len(document.tables),
    )
    blocks = read_blocks(document)
    category_blocks, reference_blocks = split_document(blocks, report)
    source_catalog, source_slugs_by_category = parse_references(
        reference_blocks, report
    )
    category_files: dict[str, dict[str, Any]] = {}

    for display_order, category_slug in enumerate(REQUIRED_CATEGORY_SLUGS, start=1):
        blocks_for_category = category_blocks.get(category_slug, [])
        if not blocks_for_category:
            continue
        description = first_general_paragraph(blocks_for_category)
        if description is None:
            report.errors.append(
                f"{category_slug}: deskripsi kategori tidak ditemukan sebelum bagian terstruktur."
            )
            description = ""
        body = blocks_to_markdown(blocks_for_category)
        if not body:
            report.errors.append(f"{category_slug}: body materi kosong.")
        professions = extract_professions(category_slug, blocks_for_category, report)
        if not professions:
            report.errors.append(f"{category_slug}: profesi tidak ditemukan.")
        profile_items = extract_profile_items(
            category_slug, blocks_for_category, report
        )
        profile_sources = [
            {
                "sourceSlug": source_slug,
                "citationLabel": None,
                "pageReference": None,
                "displayOrder": index,
            }
            for index, source_slug in enumerate(
                source_slugs_by_category.get(category_slug, [])
            )
        ]
        category_name = CATEGORY_NAMES[category_slug]
        category_files[category_slug] = {
            "schemaVersion": SCHEMA_VERSION,
            "categorySlug": category_slug,
            "category": {
                "name": category_name,
                "slug": category_slug,
                "description": description,
                "displayOrder": display_order,
                "isActive": True,
            },
            "profile": {
                "title": category_name,
                "slug": category_slug,
                "excerpt": description,
                "body": body,
                "coverImageUrl": None,
                "status": "published",
                "publishedAt": published_at,
                "readingTimeMinutes": None,
                "isFeatured": False,
                "metadata": {},
                "sources": profile_sources,
            },
            "professions": professions,
            "profileItems": profile_items,
        }

    manifest = {
        "schemaVersion": SCHEMA_VERSION,
        "datasetName": "MineVision Career",
        "description": f"Data karier yang diekstrak dari {input_path.name}.",
        "effectiveDate": effective_date,
        "categoryFiles": [
            {
                "categorySlug": slug,
                "filePath": f"categories/{slug}.json",
            }
            for slug in REQUIRED_CATEGORY_SLUGS
        ],
        "sourceCatalog": source_catalog,
    }
    validate_generated_dataset(manifest, category_files, report)
    return manifest, category_files, report


def validate_generated_dataset(
    manifest: dict[str, Any],
    category_files: dict[str, dict[str, Any]],
    report: ParseReport,
) -> None:
    manifest_slugs = [entry["categorySlug"] for entry in manifest["categoryFiles"]]
    if manifest_slugs != list(REQUIRED_CATEGORY_SLUGS):
        report.errors.append("Manifest tidak memuat 13 categorySlug wajib secara berurutan.")
    if len(category_files) != 13:
        report.errors.append(
            f"Hasil ekstraksi harus berisi 13 kategori, ditemukan {len(category_files)}."
        )
    source_slugs = {source["slug"] for source in manifest["sourceCatalog"]}
    if not source_slugs:
        report.errors.append("Source catalog kosong.")
    category_orders: list[int] = []

    for category_slug in REQUIRED_CATEGORY_SLUGS:
        career_file = category_files.get(category_slug)
        if career_file is None:
            continue
        if not (
            career_file["categorySlug"]
            == career_file["category"]["slug"]
            == career_file["profile"]["slug"]
        ):
            report.errors.append(f"{category_slug}: slug hasil ekstraksi tidak konsisten.")
        category_orders.append(career_file["category"]["displayOrder"])
        if not career_file["professions"]:
            report.errors.append(f"{category_slug}: professions kosong.")
        if not career_file["profileItems"]:
            report.errors.append(f"{category_slug}: profileItems kosong.")
        sections = {item["section"] for item in career_file["profileItems"]}
        missing_sections = set(REQUIRED_PROFILE_SECTIONS) - sections
        if missing_sections:
            report.errors.append(
                f"{category_slug}: section wajib hilang: {', '.join(sorted(missing_sections))}."
            )
        referenced_sources = {
            reference["sourceSlug"] for reference in career_file["profile"]["sources"]
        }
        unknown_sources = referenced_sources - source_slugs
        if unknown_sources:
            report.errors.append(
                f"{category_slug}: source tidak terdapat dalam catalog: {', '.join(sorted(unknown_sources))}."
            )
        if not referenced_sources:
            report.errors.append(f"{category_slug}: profile.sources kosong.")

        profession_keys = [
            (profession["groupKey"], profession["slug"])
            for profession in career_file["professions"]
        ]
        if len(profession_keys) != len(set(profession_keys)):
            report.errors.append(f"{category_slug}: natural key profesi duplikat.")
        item_keys = [item["itemKey"] for item in career_file["profileItems"]]
        if len(item_keys) != len(set(item_keys)):
            report.errors.append(f"{category_slug}: itemKey profile item duplikat.")

    if sorted(category_orders) != list(range(1, 14)):
        report.errors.append("Display order kategori wajib mencakup angka 1 sampai 13.")


def atomic_write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        dir=path.parent,
        prefix=f".{path.name}.",
        suffix=".tmp",
        text=True,
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(data, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_path, path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise


def write_dataset(
    output_path: Path,
    manifest: dict[str, Any],
    category_files: dict[str, dict[str, Any]],
    force: bool,
) -> None:
    targets = [output_path / "manifest.json"] + [
        output_path / "categories" / f"{slug}.json"
        for slug in REQUIRED_CATEGORY_SLUGS
    ]
    existing_targets = [path for path in targets if path.exists()]
    if existing_targets and not force:
        formatted = "\n".join(f"- {path}" for path in existing_targets)
        raise FileExistsError(
            "File output sudah tersedia. Gunakan --force untuk menimpa:\n" + formatted
        )
    atomic_write_json(output_path / "manifest.json", manifest)
    for category_slug in REQUIRED_CATEGORY_SLUGS:
        atomic_write_json(
            output_path / "categories" / f"{category_slug}.json",
            category_files[category_slug],
        )


def parse_iso_date(value: str) -> str:
    try:
        parsed = date.fromisoformat(value)
    except ValueError as error:
        raise argparse.ArgumentTypeError(
            "Tanggal efektif harus menggunakan format YYYY-MM-DD"
        ) from error
    if parsed.isoformat() != value:
        raise argparse.ArgumentTypeError(
            "Tanggal efektif harus menggunakan format YYYY-MM-DD"
        )
    return value


def parse_iso_datetime(value: str) -> str:
    candidate = value[:-1] + "+00:00" if value.endswith("Z") else value
    try:
        parsed = datetime.fromisoformat(candidate)
    except ValueError as error:
        raise argparse.ArgumentTypeError(
            "Published at harus berupa ISO datetime dengan timezone"
        ) from error
    if parsed.tzinfo is None:
        raise argparse.ArgumentTypeError(
            "Published at harus berupa ISO datetime dengan timezone"
        )
    return value


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Ekstrak dokumen Career MineVision menjadi kontrak staging JSON."
    )
    parser.add_argument("--input", required=True, type=Path, help="Lokasi DOCX sumber")
    parser.add_argument(
        "--output", required=True, type=Path, help="Folder output staging Career"
    )
    parser.add_argument(
        "--effective-date", required=True, type=parse_iso_date, help="Tanggal YYYY-MM-DD"
    )
    parser.add_argument(
        "--published-at",
        required=True,
        type=parse_iso_datetime,
        help="ISO datetime dengan timezone",
    )
    parser.add_argument(
        "--check-only",
        action="store_true",
        help="Parse dan validasi tanpa menulis folder atau file",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Izinkan penimpaan file output yang sudah tersedia",
    )
    return parser


def print_report(
    report: ParseReport,
    category_files: dict[str, dict[str, Any]],
    manifest: dict[str, Any],
    check_only: bool,
) -> None:
    profession_count = sum(
        len(career_file["professions"]) for career_file in category_files.values()
    )
    profile_item_count = sum(
        len(career_file["profileItems"]) for career_file in category_files.values()
    )
    section_totals = {
        section: sum(
            1
            for career_file in category_files.values()
            for item in career_file["profileItems"]
            if item["section"] == section
        )
        for section in REQUIRED_PROFILE_SECTIONS
    }
    print("\nLaporan ekstraksi Career")
    print(f"Mode          : {'check-only' if check_only else 'write'}")
    print(f"Paragraph     : {report.paragraph_count}")
    print(f"Table         : {report.table_count}")
    print(f"Categories    : {len(category_files)}")
    print(f"Professions   : {profession_count}")
    print(f"Profile item  : {profile_item_count}")
    print(f"Software table: {report.software_table_count}")
    print(f"Sources       : {len(manifest['sourceCatalog'])}")
    print(f"Warnings      : {len(report.warnings)}")
    print(f"Errors        : {len(report.errors)}")

    if category_files:
        print("\nJumlah aktual per kategori:")
        print(
            "Kategori | Profesi | Ruang lingkup | Kompetensi | Pendidikan | Software | Pelatihan | Tabel software | Sumber"
        )
        for category_slug in REQUIRED_CATEGORY_SLUGS:
            career_file = category_files.get(category_slug)
            if career_file is None:
                continue
            counts = {
                section: sum(
                    1
                    for item in career_file["profileItems"]
                    if item["section"] == section
                )
                for section in REQUIRED_PROFILE_SECTIONS
            }
            print(
                f"{category_slug} | {len(career_file['professions'])} | "
                f"{counts['work_scope']} | {counts['competency']} | "
                f"{counts['education']} | {counts['software']} | "
                f"{counts['training']} | "
                f"{report.software_tables_by_category.get(category_slug, 0)} | "
                f"{len(career_file['profile']['sources'])}"
            )

        print("\nTotal lima jenis profile item:")
        for section in REQUIRED_PROFILE_SECTIONS:
            print(f"- {section}: {section_totals[section]}")

    if report.warnings:
        print("\nWarning:")
        for warning in report.warnings:
            print(f"- {warning}")
    if report.errors:
        print("\nError struktur:", file=sys.stderr)
        for error in report.errors:
            print(f"- {error}", file=sys.stderr)
    elif check_only:
        print("\nStruktur dokumen valid. Tidak ada folder atau file output yang dibuat.")


def main() -> int:
    arguments = build_argument_parser().parse_args()
    input_path = arguments.input.resolve()
    output_path = arguments.output.resolve()

    try:
        input_status = input_path.lstat()
    except FileNotFoundError:
        print(f"File DOCX tidak ditemukan: {input_path}", file=sys.stderr)
        return 1
    if input_path.is_symlink() or not stat.S_ISREG(input_status.st_mode):
        print("Input harus berupa regular file DOCX dan bukan symbolic link.", file=sys.stderr)
        return 1
    if input_path.suffix.casefold() != ".docx":
        print("Input wajib menggunakan ekstensi .docx.", file=sys.stderr)
        return 1

    try:
        document = Document(input_path)
    except PermissionError:
        print(
            "Dokumen tidak dapat dibaca karena sedang dikunci aplikasi lain. Tutup dokumen di Microsoft Word lalu coba kembali.",
            file=sys.stderr,
        )
        return 1
    except Exception as error:
        print(f"Dokumen DOCX tidak dapat dibaca: {error}", file=sys.stderr)
        return 1

    manifest, category_files, report = build_dataset(
        document,
        input_path,
        arguments.effective_date,
        arguments.published_at,
    )
    print_report(report, category_files, manifest, arguments.check_only)
    if report.errors:
        return 1
    if arguments.check_only:
        return 0

    try:
        write_dataset(
            output_path,
            manifest,
            category_files,
            arguments.force,
        )
    except Exception as error:
        print(f"Output staging tidak dapat ditulis: {error}", file=sys.stderr)
        return 1
    print(f"\nDataset Career ditulis secara atomic ke: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
